# -*- coding: utf-8 -*-
"""读回真实配方导出的 xlsx，确认 Excel 可打开且内容正确。"""
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(HERE, 'recipe-export.xlsx')

if not os.path.exists(TARGET):
    print(f'FAIL: 未找到 {TARGET}')
    sys.exit(1)

workbook = openpyxl.load_workbook(TARGET)
print(f'openpyxl 打开成功，工作表：{workbook.sheetnames}')

failures = []
expected_sheets = ['配方信息', '原料用量', '目标达成', '形态核对']
if workbook.sheetnames != expected_sheets:
    failures.append(f'工作表不符：{workbook.sheetnames}')

info = workbook['配方信息']
info_rows = {str(r[0]): r[1] for r in info.iter_rows(values_only=True) if r[0]}
if info_rows.get('作物') != '番茄':
    failures.append(f"作物不符：{info_rows.get('作物')}")
else:
    print(f"  配方信息：{info_rows.get('作物')} / {info_rows.get('生育期')} / {info_rows.get('配制体积')}")

usage = workbook['原料用量']
usage_rows = list(usage.iter_rows(values_only=True))
print(f'  原料用量表：{len(usage_rows) - 2} 种原料 + 表头 + 合计')

# 数字必须是数值类型，才能继续在 Excel 里计算
sample = usage_rows[1]
if not isinstance(sample[2], (int, float)):
    failures.append(f'用量不是数值类型：{sample[2]!r} ({type(sample[2]).__name__})')
else:
    print(f'  用量数值类型正确：{sample[0]} → {sample[2]} g')

# 化学式应带 Unicode 下标
formulas = [row[1] for row in usage_rows[1:-1] if row[1]]
has_subscript = any(any('\u2080' <= ch <= '\u2089' for ch in str(f)) for f in formulas)
if not has_subscript:
    failures.append('化学式缺少 Unicode 下标')
else:
    print(f'  化学式下标示例：{next(f for f in formulas if any(chr(0x2080) <= ch <= chr(0x2089) for ch in str(f)))}')

# 合计行应等于各行之和
body = [row[2] for row in usage_rows[1:-1] if isinstance(row[2], (int, float))]
total_row = usage_rows[-1][2]
if isinstance(total_row, (int, float)) and abs(sum(body) - total_row) > 0.01:
    failures.append(f'合计不符：{sum(body):.3f} != {total_row}')
else:
    print(f'  合计校验通过：{total_row} g')

achieve = workbook['目标达成']
achieve_rows = list(achieve.iter_rows(values_only=True))
print(f'  目标达成表：{len(achieve_rows) - 1} 个元素')

if failures:
    print('\nFAIL:')
    for item in failures:
        print(f'  - {item}')
    sys.exit(1)

print('\nPASS: 真实配方导出正确，Excel 可正常读取')
sys.exit(0)
