# -*- coding: utf-8 -*-
"""
用 openpyxl（独立第三方实现）读取 Node 生成的 xlsx，验证文件结构合法。

为什么要交叉验证：xlsx 是 ZIP + XML，自己写 ZIP 最容易在 CRC32、
中央目录偏移这些地方出错，而**自测永远发现不了**（自己写的读自己写的）。
用另一个语言的成熟库打开一次，才算真的验证过。
"""
import os
import sys

try:
    import openpyxl
except ImportError:
    print('SKIP: 未安装 openpyxl')
    sys.exit(0)

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(HERE, 'xlsx-check-output.xlsx')

if not os.path.exists(TARGET):
    print(f'FAIL: 文件不存在 {TARGET}')
    sys.exit(1)

failures = []

try:
    workbook = openpyxl.load_workbook(TARGET)
except Exception as error:  # noqa: BLE001
    print(f'FAIL: openpyxl 无法打开文件 —— {error}')
    sys.exit(1)

print(f'openpyxl 成功打开，工作表：{workbook.sheetnames}')

expected_sheets = ['原料用量', '目标达成', '中文与特殊字符']
if workbook.sheetnames != expected_sheets:
    failures.append(f'工作表名不符：{workbook.sheetnames} != {expected_sheets}')

sheet1 = workbook['原料用量']
rows = list(sheet1.iter_rows(values_only=True))
if rows[0][0] != '原料':
    failures.append(f'表头不符：{rows[0]}')
if rows[1][1] != 118.0:
    failures.append(f'数字未被识别为数值：{rows[1]}')
else:
    print(f'  数字类型正确：{rows[1][1]!r}（{type(rows[1][1]).__name__}）')
if 'Ca(NO₃)₂·4H₂O' not in str(rows[1][0]):
    failures.append(f'Unicode 下标丢失：{rows[1][0]}')
else:
    print(f'  Unicode 下标保留：{rows[1][0]}')
if rows[-1][0] != '合计':
    failures.append(f'最后一行不符：{rows[-1]}')

sheet3 = workbook['中文与特殊字符']
rows3 = list(sheet3.iter_rows(values_only=True))
if rows3[0][0] != '含 & < > " 的文本':
    failures.append(f'XML 转义有误：{rows3[0][0]!r}')
else:
    print(f'  XML 特殊字符正确还原：{rows3[0][0]!r}')
if rows3[2][1] != -12.5:
    failures.append(f'负数不正确：{rows3[2][1]}')

# 检查表头加粗是否生效
font = sheet1.cell(row=1, column=1).font
if not font.bold:
    failures.append('表头加粗未生效')
else:
    print('  表头加粗已生效')

# 检查列宽是否写入
width = sheet1.column_dimensions['A'].width
print(f'  A 列宽度：{width}')

if failures:
    print('\nFAIL:')
    for item in failures:
        print(f'  - {item}')
    sys.exit(1)

print('\nPASS: 生成的文件结构合法，内容与数值类型全部正确')
sys.exit(0)
